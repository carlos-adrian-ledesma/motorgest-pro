from __future__ import annotations

import base64
import io
from pathlib import Path
from typing import Any

from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether,
)

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

NAVY = colors.HexColor('#0B355D')
BLUE = colors.HexColor('#1D89D8')
SILVER = colors.HexColor('#E7EDF2')
MID = colors.HexColor('#9AA8B5')
DARK = colors.HexColor('#172431')
TEXT = colors.HexColor('#1D2A35')


def _image_bytes(data_url: str) -> bytes:
    if not data_url:
        raise ValueError('La imagen de inspección está vacía.')
    raw = data_url.split(',', 1)[1] if ',' in data_url else data_url
    return base64.b64decode(raw)


def _safe(value: Any, fallback: str = '—') -> str:
    text = str(value or '').strip()
    return text or fallback


def _fit_image_size(image_bytes: bytes, max_w_mm: float, max_h_mm: float) -> tuple[float, float]:
    with PILImage.open(io.BytesIO(image_bytes)) as im:
        w, h = im.size
    scale = min((max_w_mm * mm) / w, (max_h_mm * mm) / h)
    return w * scale, h * scale


def _pdf_footer(canvas, doc, version: str):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor('#D3DCE4'))
    canvas.line(14 * mm, 12 * mm, 196 * mm, 12 * mm)
    canvas.setFillColor(colors.HexColor('#607080'))
    canvas.setFont('Helvetica', 7.2)
    canvas.drawString(14 * mm, 7.5 * mm, f'MotorGest Pro V{version} · Documento generado localmente')
    canvas.drawRightString(196 * mm, 7.5 * mm, f'Página {doc.page}')
    canvas.restoreState()


def build_inspection_pdf(payload: dict[str, Any], output_path: Path, logo_path: Path, version: str) -> Path:
    output_path = Path(output_path)
    image_bytes = _image_bytes(payload.get('imageDataUrl', ''))
    doc = SimpleDocTemplate(
        str(output_path), pagesize=A4,
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=10 * mm, bottomMargin=16 * mm,
        title=f"Ficha de inspección visual - {_safe(payload.get('caseNo'))}",
        author='MotorGest Pro',
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('MGTitle', parent=styles['Title'], fontName='Helvetica-Bold',
                                 fontSize=17, leading=19, textColor=NAVY, alignment=TA_RIGHT, spaceAfter=1)
    sub_style = ParagraphStyle('MGSub', parent=styles['Normal'], fontName='Helvetica-Bold',
                               fontSize=9.2, leading=11, textColor=BLUE, alignment=TA_RIGHT)
    label_style = ParagraphStyle('MGLabel', parent=styles['Normal'], fontName='Helvetica-Bold',
                                 fontSize=7.8, leading=9.5, textColor=NAVY)
    value_style = ParagraphStyle('MGValue', parent=styles['Normal'], fontName='Helvetica',
                                 fontSize=8.2, leading=9.8, textColor=TEXT)
    section_style = ParagraphStyle('MGSection', parent=styles['Heading2'], fontName='Helvetica-Bold',
                                   fontSize=10.5, leading=12, textColor=NAVY, spaceBefore=2, spaceAfter=4)
    note_style = ParagraphStyle('MGNote', parent=styles['Normal'], fontName='Helvetica',
                                fontSize=8.2, leading=11, textColor=TEXT)

    story = []
    logo = None
    if logo_path.exists():
        logo = Image(str(logo_path), width=24*mm, height=24*mm)
    title_box = [Paragraph('FICHA DE INSPECCIÓN VISUAL', title_style),
                 Paragraph(_safe(payload.get('unitType'), 'Unidad'), sub_style),
                 Paragraph(_safe(payload.get('company'), 'MotorGest Pro'), ParagraphStyle(
                     'Company', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5,
                     textColor=colors.HexColor('#6C7B88'), alignment=TA_RIGHT))]
    header = Table([[logo if logo else '', title_box]], colWidths=[30*mm, 152*mm], rowHeights=[27*mm])
    header.setStyle(TableStyle([
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('ALIGN',(0,0),(0,0),'LEFT'),
        ('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F7F9FB')),
        ('BOX',(0,0),(-1,-1),0.7,colors.HexColor('#AAB8C4')),
        ('LINEBELOW',(0,0),(-1,-1),2.2,NAVY),
        ('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),
        ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
    ]))
    story += [header, Spacer(1, 4*mm)]

    meta_rows = [
        [Paragraph('Cliente', label_style), Paragraph(_safe(payload.get('client')), value_style),
         Paragraph('Caso / OT', label_style), Paragraph(_safe(payload.get('caseNo')), value_style)],
        [Paragraph('Unidad', label_style), Paragraph(_safe(payload.get('unitType')), value_style),
         Paragraph('Dominio / ID', label_style), Paragraph(_safe(payload.get('domainId')), value_style)],
        [Paragraph('Marca / Modelo', label_style), Paragraph(_safe(payload.get('brandModel')), value_style),
         Paragraph('Año', label_style), Paragraph(_safe(payload.get('year')), value_style)],
        [Paragraph('Km / Horas', label_style), Paragraph(_safe(payload.get('kmHours')), value_style),
         Paragraph('Fecha / Inspector', label_style), Paragraph(f"{_safe(payload.get('date'))} · {_safe(payload.get('inspector'))}", value_style)],
    ]
    meta = Table(meta_rows, colWidths=[24*mm, 67*mm, 25*mm, 66*mm])
    meta.setStyle(TableStyle([
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('BACKGROUND',(0,0),(0,-1),SILVER),('BACKGROUND',(2,0),(2,-1),SILVER),
        ('GRID',(0,0),(-1,-1),0.45,colors.HexColor('#AAB8C4')),
        ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
        ('TOPPADDING',(0,0),(-1,-1),3.5),('BOTTOMPADDING',(0,0),(-1,-1),3.5),
    ]))
    story += [meta, Spacer(1, 4*mm)]

    iw, ih = _fit_image_size(image_bytes, 182, 91)
    inspection_img = Image(io.BytesIO(image_bytes), width=iw, height=ih)
    img_frame = Table([[inspection_img]], colWidths=[182*mm])
    img_frame.setStyle(TableStyle([
        ('BOX',(0,0),(-1,-1),0.8,colors.HexColor('#7E91A1')),
        ('LEFTPADDING',(0,0),(-1,-1),1),('RIGHTPADDING',(0,0),(-1,-1),1),
        ('TOPPADDING',(0,0),(-1,-1),1),('BOTTOMPADDING',(0,0),(-1,-1),1),
    ]))
    story += [img_frame, Spacer(1, 4*mm), Paragraph('REGISTRO DIGITAL DE HALLAZGOS', section_style)]

    marks = payload.get('marks') or []
    rows = [['N°','Zona / componente','Hallazgo / daño observado','Estado']]
    if marks:
        for idx, mark in enumerate(marks, 1):
            finding = ' · '.join(x for x in [str(mark.get('label') or '').strip(), str(mark.get('note') or '').strip()] if x)
            rows.append([str(idx), _safe(mark.get('zone'), ''), finding, _safe(mark.get('status'), '')])
    else:
        rows.append(['1','','',''])
    findings = Table(rows, colWidths=[10*mm, 47*mm, 91*mm, 34*mm], repeatRows=1)
    findings.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),NAVY),('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,0),7.6),
        ('FONTNAME',(0,1),(-1,-1),'Helvetica'),('FONTSIZE',(0,1),(-1,-1),7.6),
        ('GRID',(0,0),(-1,-1),0.45,colors.HexColor('#AAB8C4')),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),
        ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white, colors.HexColor('#F7F9FB')]),
    ]))
    story += [findings, Spacer(1, 4*mm)]

    notes = _safe(payload.get('notes'), 'Sin observaciones generales registradas.')
    notes_table = Table([
        [Paragraph('OBSERVACIONES GENERALES', ParagraphStyle('NotesHead', parent=section_style, textColor=colors.white, spaceBefore=0, spaceAfter=0))],
        [Paragraph(notes, note_style)]
    ], colWidths=[182*mm])
    notes_table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),NAVY),('BACKGROUND',(0,1),(-1,1),colors.HexColor('#F7F9FB')),
        ('BOX',(0,0),(-1,-1),0.6,colors.HexColor('#AAB8C4')),
        ('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),
        ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),5),
    ]))
    sign = Table([
        ['',''],
        [Paragraph('Firma inspector', ParagraphStyle('Sig', parent=value_style, alignment=TA_CENTER)),
         Paragraph('Firma cliente', ParagraphStyle('Sig2', parent=value_style, alignment=TA_CENTER))]
    ], colWidths=[91*mm,91*mm], rowHeights=[11*mm,6*mm])
    sign.setStyle(TableStyle([
        ('LINEABOVE',(0,1),(0,1),0.7,NAVY),('LINEABOVE',(1,1),(1,1),0.7,NAVY),
        ('LEFTPADDING',(0,0),(-1,-1),10),('RIGHTPADDING',(0,0),(-1,-1),10),
        ('TOPPADDING',(0,0),(-1,-1),2),('BOTTOMPADDING',(0,0),(-1,-1),2),
    ]))
    story += [notes_table, Spacer(1, 6*mm), KeepTogether(sign)]

    doc.build(story, onFirstPage=lambda c,d:_pdf_footer(c,d,version), onLaterPages=lambda c,d:_pdf_footer(c,d,version))
    return output_path


def _shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def _set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')


def _set_cell_text(cell, text: str, bold=False, color='1D2A35', size=8.5, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold; r.font.name = 'Arial'; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge, spec in edges.items():
        tag = f'w:{edge}'
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn('w:val'), spec.get('val', 'single'))
        node.set(qn('w:sz'), str(spec.get('sz', 6)))
        node.set(qn('w:color'), spec.get('color', 'AAB8C4'))


def _section_band(document, text: str):
    t = document.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    _shade_cell(c, '0B355D')
    _set_cell_text(c, text, bold=True, color='FFFFFF', size=8.7)
    c.paragraphs[0].paragraph_format.space_before = Pt(0)
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    return t


def build_inspection_docx(payload: dict[str, Any], output_path: Path, logo_path: Path, version: str) -> Path:
    output_path = Path(output_path)
    image_bytes = _image_bytes(payload.get('imageDataUrl', ''))
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210); section.page_height = Mm(297)
    section.top_margin = Mm(8); section.bottom_margin = Mm(10); section.left_margin = Mm(10); section.right_margin = Mm(10)
    section.header_distance = Mm(4); section.footer_distance = Mm(5)

    styles = document.styles
    styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(8.5)
    styles['Normal'].paragraph_format.space_after = Pt(0)

    # Cabecera institucional de tres zonas.
    header = document.add_table(rows=1, cols=3)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER; header.autofit = False
    widths = [Mm(31), Mm(101), Mm(56)]
    for i, width in enumerate(widths): header.columns[i].width = width
    c0, c1, c2 = header.rows[0].cells
    for c in (c0, c1, c2):
        _shade_cell(c, 'F3F6F9'); _set_cell_margins(c, top=70, start=90, bottom=70, end=90)
        _cell_border(c, bottom={'sz': 14, 'color': '0B355D'})
    if logo_path.exists():
        p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run().add_picture(str(logo_path), width=Mm(24))
    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p1.add_run('FICHA DE INSPECCIÓN VISUAL'); r.bold = True; r.font.name = 'Aptos Display'; r.font.size = Pt(19); r.font.color.rgb = RGBColor(11,53,93)
    p1.paragraph_format.space_after = Pt(1)
    p1b = c1.add_paragraph(); p1b.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p1b.add_run('CONTROL DIGITAL · RECEPCIÓN · TRAZABILIDAD'); r.bold = True; r.font.size = Pt(7.3); r.font.color.rgb = RGBColor(92,111,127)
    p1c = c1.add_paragraph(); p1c.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p1c.add_run(_safe(payload.get('company'), 'MotorGest Pro')); r.font.size = Pt(8); r.font.color.rgb = RGBColor(29,137,216)
    _shade_cell(c2, 'E7EDF2')
    p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p2.add_run(_safe(payload.get('unitType'), 'UNIDAD').upper()); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = RGBColor(11,53,93)
    p2b = c2.add_paragraph(); p2b.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p2b.add_run(f"CASO / OT  {_safe(payload.get('caseNo'))}"); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(29,137,216)
    p2c = c2.add_paragraph(); p2c.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p2c.add_run(f"FECHA  {_safe(payload.get('date'))}"); r.font.size = Pt(7.6); r.font.color.rgb = RGBColor(75,92,108)

    spacer = document.add_paragraph(); spacer.paragraph_format.space_after = Pt(2); spacer.paragraph_format.space_before = Pt(0)

    _section_band(document, 'DATOS GENERALES DE LA UNIDAD')
    meta = document.add_table(rows=4, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER; meta.autofit = False
    widths=[Mm(25),Mm(69),Mm(27),Mm(67)]
    for row in meta.rows:
        for i,w in enumerate(widths): row.cells[i].width=w
    meta_values = [
        ('Cliente', _safe(payload.get('client')), 'Dominio / ID', _safe(payload.get('domainId'))),
        ('Marca / Modelo', _safe(payload.get('brandModel')), 'Año', _safe(payload.get('year'))),
        ('Km / Horas', _safe(payload.get('kmHours')), 'N° serie / VIN', _safe(payload.get('serial'))),
        ('Inspector', _safe(payload.get('inspector')), 'Fecha', _safe(payload.get('date'))),
    ]
    for r_idx, vals in enumerate(meta_values):
        for c_idx, text in enumerate(vals):
            label = c_idx in (0,2)
            cell = meta.cell(r_idx,c_idx)
            _set_cell_text(cell, text, bold=label, color='0B355D' if label else '1D2A35', size=8.0)
            if label: _shade_cell(cell,'E7EDF2')
            else: _shade_cell(cell,'FBFCFD')
            _cell_border(cell, top={'sz':4,'color':'BCC7D1'}, bottom={'sz':4,'color':'BCC7D1'}, start={'sz':4,'color':'BCC7D1'}, end={'sz':4,'color':'BCC7D1'})

    # Estado operativo visible y separado de los datos identificatorios.
    state_table = document.add_table(rows=1, cols=3)
    state_table.alignment = WD_TABLE_ALIGNMENT.CENTER; state_table.autofit = False
    state_labels = [
        ('ESTADO DEL CASO', _safe(payload.get('caseStatus'), 'Ingresado')),
        ('ESTADO DE INSPECCIÓN', _safe(payload.get('generalState'), 'No evaluado')),
        ('TIPO DE UNIDAD', _safe(payload.get('unitType'), 'Unidad')),
    ]
    for idx,(lab,val) in enumerate(state_labels):
        cell=state_table.cell(0,idx); cell.width=Mm(62.6); _shade_cell(cell,'F3F7FA'); _set_cell_margins(cell, top=70,start=110,bottom=70,end=110)
        p0=cell.paragraphs[0]; p0.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=p0.add_run(lab); rr.bold=True; rr.font.size=Pt(6.8); rr.font.color.rgb=RGBColor(92,111,127)
        p1=cell.add_paragraph(); p1.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=p1.add_run(val); rr.bold=True; rr.font.size=Pt(9.2); rr.font.color.rgb=RGBColor(11,53,93)
        _cell_border(cell, bottom={'sz':5,'color':'1D89D8'}, top={'sz':3,'color':'D7E0E8'}, start={'sz':3,'color':'D7E0E8'}, end={'sz':3,'color':'D7E0E8'})

    p_gap = document.add_paragraph(); p_gap.paragraph_format.space_after=Pt(1); p_gap.paragraph_format.space_before=Pt(0)
    _section_band(document, 'INSPECCIÓN VISUAL DIGITAL')
    p_img = document.add_paragraph(); p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before=Pt(2); p_img.paragraph_format.space_after=Pt(2)
    p_img.add_run().add_picture(io.BytesIO(image_bytes), width=Mm(188))
    p_cap=document.add_paragraph(); p_cap.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p_cap.paragraph_format.space_after=Pt(2)
    rr=p_cap.add_run(f"Marcas digitales registradas: {len(payload.get('marks') or [])}"); rr.bold=True; rr.font.size=Pt(7); rr.font.color.rgb=RGBColor(92,111,127)

    _section_band(document, 'REGISTRO DIGITAL DE HALLAZGOS')
    marks = payload.get('marks') or []
    table = document.add_table(rows=1 + max(1,len(marks)), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'; table.autofit=False
    tw=[Mm(10),Mm(46),Mm(96),Mm(36)]
    headers=['N°','Zona / componente','Hallazgo / daño observado','Estado']
    for i,h in enumerate(headers):
        table.columns[i].width=tw[i]; _set_cell_text(table.cell(0,i),h,bold=True,color='FFFFFF',size=7.4,align=WD_ALIGN_PARAGRAPH.CENTER); _shade_cell(table.cell(0,i),'164D78')
    data = marks if marks else [{}]
    for ridx, mark in enumerate(data,1):
        finding=' · '.join(x for x in [str(mark.get('label') or '').strip(),str(mark.get('note') or '').strip()] if x)
        vals=[str(ridx),str(mark.get('zone') or ''),finding,str(mark.get('status') or '')]
        for cidx,val in enumerate(vals):
            cell=table.cell(ridx,cidx); _set_cell_text(cell,val,size=7.5,align=WD_ALIGN_PARAGRAPH.CENTER if cidx in (0,3) else WD_ALIGN_PARAGRAPH.LEFT)
            _shade_cell(cell,'F8FAFC' if ridx%2 else 'EEF3F7')

    gap=document.add_paragraph(); gap.paragraph_format.space_after=Pt(1); gap.paragraph_format.space_before=Pt(0)
    _section_band(document, 'OBSERVACIONES Y CONFORMIDAD')
    notes_sign = document.add_table(rows=1, cols=2)
    notes_sign.alignment=WD_TABLE_ALIGNMENT.CENTER; notes_sign.autofit=False
    left,right=notes_sign.rows[0].cells; left.width=Mm(116); right.width=Mm(72)
    _shade_cell(left,'F8FAFC'); _shade_cell(right,'F3F6F9')
    _set_cell_margins(left,top=100,start=120,bottom=100,end=120); _set_cell_margins(right,top=100,start=120,bottom=100,end=120)
    p=left.paragraphs[0]; r=p.add_run(_safe(payload.get('notes'),'Sin observaciones generales registradas.')); r.font.size=Pt(8.2); r.font.color.rgb=RGBColor(29,42,53)
    p.paragraph_format.space_after=Pt(2)
    p=right.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('CONFORMIDAD'); r.bold=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor(11,53,93)
    p=right.add_paragraph('\n'); p.paragraph_format.space_after=Pt(0)
    p=right.add_paragraph('Firma inspector'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size=Pt(7.2); p.runs[0].font.color.rgb=RGBColor(75,92,108)
    p=right.add_paragraph('\n'); p.paragraph_format.space_after=Pt(0)
    p=right.add_paragraph('Firma cliente'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size=Pt(7.2); p.runs[0].font.color.rgb=RGBColor(75,92,108)
    _cell_border(left, top={'sz':4,'color':'BCC7D1'}, bottom={'sz':4,'color':'BCC7D1'}, start={'sz':4,'color':'BCC7D1'}, end={'sz':4,'color':'BCC7D1'})
    _cell_border(right, top={'sz':4,'color':'BCC7D1'}, bottom={'sz':4,'color':'BCC7D1'}, start={'sz':4,'color':'BCC7D1'}, end={'sz':4,'color':'BCC7D1'})

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr=footer.add_run(f'MotorGest Pro V{version}  |  Documento Word nativo  |  {_safe(payload.get("caseNo"))}')
    rr.font.name='Aptos'; rr.font.size=Pt(6.8); rr.font.color.rgb=RGBColor(96,112,128)

    props=document.core_properties
    props.title=f"Ficha de inspección visual - {_safe(payload.get('caseNo'))}"
    props.subject='Inspección visual digital de vehículo'
    props.author='MotorGest Pro'
    props.keywords='MotorGest Pro, inspección visual, taller, vehículo'
    document.save(str(output_path))
    return output_path
